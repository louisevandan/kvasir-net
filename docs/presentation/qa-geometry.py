# 기하 QA: 슬라이드 경계 이탈, 텍스트 넘침 추정, 표 높이 추정
import sys, math, unicodedata
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
W, H = 13.3333, 7.5

def adv(ch, mono):
    if ch == '\n':
        return 0.0
    e = unicodedata.east_asian_width(ch)
    if e in ('W', 'F'):
        return 1.0
    if mono:
        return 0.6
    if ch in ' .,:·|()[]{}ilj!\'"':
        return 0.30
    if ch.isupper():
        return 0.62
    return 0.52

def text_w_pt(txt, size, mono):
    return sum(adv(c, mono) for c in txt) * size

def est_lines(txt, box_w_pt, size, mono):
    n = 0
    for para in txt.split('\n'):
        if not para:
            n += 1
            continue
        w = text_w_pt(para, size, mono)
        n += max(1, math.ceil(w / max(box_w_pt, 1)))
    return n

def run_info(tf):
    """가장 큰 글자 크기와 mono 여부, 전체 텍스트(문단 줄바꿈 유지)"""
    sizes = []
    mono = False
    parts = []
    for p in tf.paragraphs:
        line = ''.join(r.text for r in p.runs)
        parts.append(line)
        for r in p.runs:
            if r.font.size:
                sizes.append(r.font.size.pt)
            if r.font.name and 'Mono' in r.font.name:
                mono = True
    size = max(sizes) if sizes else 10.5
    return '\n'.join(parts), size, mono

prs = Presentation(sys.argv[1])
issues = []
for i, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        try:
            x, y = sh.left / EMU_IN, sh.top / EMU_IN
            w, h = sh.width / EMU_IN, sh.height / EMU_IN
        except TypeError:
            continue
        if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
            label = (sh.text_frame.text[:40].replace('\n', ' ') if sh.has_text_frame else sh.shape_type)
            issues.append(f'S{i:02d} OUT-OF-BOUNDS  x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}  {label}')
        if sh.has_table:
            tbl = sh.table
            th = 0
            for r_i, row in enumerate(tbl.rows):
                cell_lines = 1
                for c_i, cell in enumerate(row.cells):
                    txt, size, mono = run_info(cell.text_frame)
                    cw = tbl.columns[c_i].width / EMU_IN * 72 - 10
                    cell_lines = max(cell_lines, est_lines(txt, cw, size, mono))
                th += cell_lines * size * 1.25 + 10
            th_in = th / 72
            if y + th_in > H - 0.05:
                issues.append(f'S{i:02d} TABLE-OVERFLOW  top={y:.2f} est_h={th_in:.2f} bottom={y+th_in:.2f} > {H}')
            continue
        if not sh.has_text_frame:
            continue
        txt, size, mono = run_info(sh.text_frame)
        if not txt.strip():
            continue
        box_w_pt = w * 72
        box_h_pt = h * 72
        lines = est_lines(txt, box_w_pt, size, mono)
        need = lines * size * 1.22
        bottom_in = y + need / 72
        if bottom_in > H - 0.06:
            issues.append(f'S{i:02d} PAST-BOTTOM   y={y:.2f} est_bottom={bottom_in:.2f} | {txt[:50]}')
        if need > box_h_pt + 7:
            issues.append(
                f'S{i:02d} TEXT-OVERFLOW  need={need:.0f}pt box={box_h_pt:.0f}pt lines={lines} size={size} '
                f'| {txt[:52].replace(chr(10)," ")}')

print(f'slides={len(prs.slides)} issues={len(issues)}')
for s in issues:
    print(' ', s)
